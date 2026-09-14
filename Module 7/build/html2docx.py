"""
Convert one of the house-style answer HTMLs into an editable .docx.

The PDF is the submission artefact; the .docx is the working copy - it opens in
LibreOffice Writer and in Word, so lecturer corrections can be typed straight in.
Both come from the SAME html, so regenerate this after every edit to the html
and the two never drift apart.

Handles the subset of markup the answer sheets actually use:
  .sec bars, h3.qn (question + marks), p, ul/li, tables, .card/.note/.exam blocks,
  <b>/<strong>, <i>/<em>, <small>, <sub>, <br>.

Usage:  python html2docx.py mod7-workbook-answers.html "..\\assignment\\Module 7 - ... .docx"
"""

import html as htmllib
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

INK = RGBColor(0x16, 0x15, 0x0F)
PURPLE = RGBColor(0x5B, 0x4F, 0xC7)
MUTED = RGBColor(0x6E, 0x6D, 0x67)


# ---------------------------------------------------------------- inline runs
def runs(fragment):
    """Split an inline HTML fragment into (text, bold, italic, small) tuples."""
    out = []
    bold = ital = small = 0
    token = re.compile(r"<(/?)(b|strong|i|em|small|br|sub|sup)[^>]*>|([^<]+)", re.I)
    for m in token.finditer(fragment):
        close, tag, text = m.group(1), (m.group(2) or "").lower(), m.group(3)
        if text is not None:
            t = htmllib.unescape(re.sub(r"\s+", " ", text))
            if t:
                out.append((t, bool(bold), bool(ital), bool(small)))
            continue
        if tag == "br":
            out.append(("\n", bool(bold), bool(ital), bool(small)))
        elif tag in ("b", "strong"):
            bold += -1 if close else 1
        elif tag in ("i", "em"):
            ital += -1 if close else 1
        elif tag == "small":
            small += -1 if close else 1
    return out


def write(par, fragment, size=10.5, colour=INK):
    for text, bold, ital, small in runs(fragment):
        r = par.add_run(text)
        r.bold, r.italic = bold, ital
        r.font.size = Pt(size - 1 if small else size)
        r.font.color.rgb = MUTED if small else colour
    return par


def strip(fragment):
    return htmllib.unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", fragment))).strip()


# ---------------------------------------------------------------- block walk
def convert(html, doc):
    body = html.split("<body>", 1)[1].rsplit("</body>", 1)[0]
    # the cover sheet is PDF furniture - the docx gets a plain title instead
    sections = re.findall(r'<section class="sheet.*?</section>', body, re.S)
    sections = [s for s in sections if 'class="sheet cover"' not in s]

    seen_bars = set()
    for sec in sections:
        sec = re.sub(r'<div class="foot">.*?</div>', "", sec, flags=re.S)

        for m in re.finditer(
            r'<div class="sec">(.*?)</div>'
            r'|<h3 class="qn">(.*?)</h3>'
            r'|<h4>(.*?)</h4>'
            r'|<p[^>]*>(.*?)</p>'
            r'|<ul[^>]*>(.*?)</ul>'
            r'|<table>(.*?)</table>'
            r'|<div class="formula"[^>]*>(.*?)</div>'
            r'|<span class="t">(.*?)</span>',
            sec, re.S,
        ):
            bar, qn, h4, para, ul, table, formula, tag = m.groups()

            if bar is not None:
                label = strip(bar)
                if label in seen_bars:
                    continue
                seen_bars.add(label)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                r = p.add_run(label.upper())
                r.bold = True
                r.font.size = Pt(11)
                r.font.color.rgb = PURPLE

            elif qn is not None:
                marks = re.search(r'<span class="marks">(.*?)</span>', qn)
                text = strip(re.sub(r'<span class="marks">.*?</span>', "", qn, flags=re.S))
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(3)
                r = p.add_run(text + ("   " + strip(marks.group(1)) if marks else ""))
                r.bold = True
                r.font.size = Pt(11)

            elif h4 is not None:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(strip(h4))
                r.bold = True
                r.font.size = Pt(10.5)

            elif para is not None:
                if not strip(para):
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                write(p, para)

            elif ul is not None:
                for li in re.findall(r"<li[^>]*>(.*?)</li>", ul, re.S):
                    p = doc.add_paragraph(style="List Bullet")
                    p.paragraph_format.space_after = Pt(2)
                    write(p, li)

            elif table is not None:
                head = re.findall(r"<th[^>]*>(.*?)</th>", table, re.S)
                rows = re.findall(r"<tr>(?!.*?<th)(.*?)</tr>", table, re.S)
                rows = [re.findall(r"<td[^>]*>(.*?)</td>", r, re.S) for r in rows]
                rows = [r for r in rows if r]
                if not rows:
                    continue
                cols = max([len(head)] + [len(r) for r in rows])
                t = doc.add_table(rows=0, cols=cols)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                if head:
                    cells = t.add_row().cells
                    for i, h in enumerate(head[:cols]):
                        p = cells[i].paragraphs[0]
                        r = p.add_run(strip(h))
                        r.bold = True
                        r.font.size = Pt(9)
                        r.font.color.rgb = PURPLE
                for row in rows:
                    cells = t.add_row().cells
                    for i, c in enumerate(row[:cols]):
                        write(cells[i].paragraphs[0], c, size=9.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)

            elif formula is not None:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                write(p, formula, size=12)

            elif tag is not None:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(strip(tag).upper())
                r.bold = True
                r.font.size = Pt(8.5)
                r.font.color.rgb = PURPLE


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    src = sys.argv[1] if len(sys.argv) > 1 else "mod7-workbook-answers.html"
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        here, os.pardir, "assignment",
        "Module 7 — Foreign Exchange Markets (Workbook Answers - editable).docx")

    html = open(os.path.join(here, src), encoding="utf-8").read()
    title = strip(re.search(r"<title>(.*?)</title>", html, re.S).group(1))

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    h = doc.add_paragraph()
    r = h.add_run(title)
    r.bold = True
    r.font.size = Pt(16)
    sub = doc.add_paragraph()
    r = sub.add_run("Occupational Certificate: Financial Markets Practitioner · Novia One Business School\n"
                    "Editable working copy — type lecturer corrections straight into this file, then "
                    "mirror them back into build/" + src + " and re-run build.py to refresh the PDF.")
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED

    convert(html, doc)

    out = os.path.abspath(out)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
